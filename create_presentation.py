import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_presentation(output_path="ScamGuard_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette (ScamGuard Precision Cyber Theme)
    BG_COLOR = RGBColor(11, 15, 23)          # #0B0F17 Dark Slate/Obsidian
    CARD_BG = RGBColor(19, 27, 46)           # #131B2E Dark Navy Card
    CARD_BORDER = RGBColor(35, 50, 82)       # #233252 Subtle Border
    TEXT_WHITE = RGBColor(255, 255, 255)     # #FFFFFF
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8 Slate Gray
    TEXT_CYAN = RGBColor(56, 189, 248)       # #38BDF8 Sky Cyan Accent
    ACCENT_CYAN = RGBColor(56, 189, 248)     # #38BDF8 Sky Cyan Accent
    ACCENT_BLUE = RGBColor(37, 99, 235)      # #2563EB Primary Blue
    ACCENT_GREEN = RGBColor(16, 185, 129)    # #10B981 Emerald Safe
    ACCENT_AMBER = RGBColor(245, 158, 11)    # #F59E0B Amber Warning
    ACCENT_RED = RGBColor(239, 68, 68)       # #EF4444 Crimson Critical

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_COLOR
        bg_shape.line.fill.background()

    def add_header(slide, category, title, subtitle):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEXT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.38), Inches(11.7), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED

    def add_footer(slide, slide_num):
        # Subtle Separator Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

        # Footer Left Label
        foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(6.0), Inches(0.35))
        tf_f = foot_box.text_frame
        tf_f.margin_left = tf_f.margin_top = tf_f.margin_bottom = tf_f.margin_right = 0
        p_f = tf_f.paragraphs[0]
        p_f.text = "ScamGuard — Personal Scam Defense & Threat Intelligence Platform"
        p_f.font.name = "Segoe UI"
        p_f.font.size = Pt(9.5)
        p_f.font.color.rgb = TEXT_MUTED

        # Slide Number
        num_box = slide.shapes.add_textbox(Inches(10.5), Inches(6.95), Inches(2.033), Inches(0.35))
        tf_n = num_box.text_frame
        tf_n.margin_left = tf_n.margin_top = tf_n.margin_bottom = tf_n.margin_right = 0
        p_n = tf_n.paragraphs[0]
        p_n.alignment = PP_ALIGN.RIGHT
        p_n.text = f"{slide_num} / 8"
        p_n.font.name = "Segoe UI"
        p_n.font.size = Pt(9.5)
        p_n.font.bold = True
        p_n.font.color.rgb = TEXT_CYAN

    # =========================================================================
    # SLIDE 1: Title Slide (Executive Cover)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    # Ambient Card Container
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.1), Inches(10.333), Inches(5.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = CARD_BORDER
    card1.line.width = Pt(1.5)

    # Brand Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.766), Inches(1.5), Inches(3.8), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(15, 30, 60)
    badge.line.color.rgb = ACCENT_BLUE
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    p_b.text = "[ SCAMGUARD THREAT DEFENSE ]"
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = TEXT_CYAN

    # Main Title
    tbox = s1.shapes.add_textbox(Inches(2.0), Inches(2.1), Inches(9.333), Inches(1.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "ScamGuard"
    p.font.name = "Segoe UI"
    p.font.size = Pt(50)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Tagline & Mission
    sub = s1.shapes.add_textbox(Inches(2.0), Inches(3.25), Inches(9.333), Inches(0.8))
    tf_s = sub.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.alignment = PP_ALIGN.CENTER
    p_s.text = "Don't just detect the scam. Understand the attack."
    p_s.font.name = "Segoe UI"
    p_s.font.size = Pt(18)
    p_s.font.bold = True
    p_s.font.color.rgb = TEXT_CYAN

    p_s2 = tf_s.add_paragraph()
    p_s2.alignment = PP_ALIGN.CENTER
    p_s2.text = "Next-Generation Multi-Modal Scam Deconstruction Powered by Gemini AI"
    p_s2.font.name = "Segoe UI"
    p_s2.font.size = Pt(13)
    p_s2.font.color.rgb = TEXT_MUTED

    # 3 Highlight Feature Pills
    pills = [
        ("Multi-Modal Ingestion", "SMS, Screenshots & URLs"),
        ("Forensic Intelligence", "Step-by-step attack chains"),
        ("Zero-Retention Privacy", "Ephemeral sandbox security")
    ]
    pill_w = Inches(2.9)
    pill_gap = Inches(0.4)
    start_x = Inches(2.0)
    pill_y = Inches(4.6)

    for i, (p_title, p_sub) in enumerate(pills):
        px = start_x + i * (pill_w + pill_gap)
        p_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, pill_y, pill_w, Inches(1.1))
        p_shape.fill.solid()
        p_shape.fill.fore_color.rgb = RGBColor(13, 20, 36)
        p_shape.line.color.rgb = CARD_BORDER
        p_shape.line.width = Pt(1)

        p_tf = p_shape.text_frame
        p_tf.word_wrap = True
        p_tf.margin_top = Inches(0.18)
        p_tf.margin_left = Inches(0.15)
        p_tf.margin_right = Inches(0.15)

        p1 = p_tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = p_title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE

        p2 = p_tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = p_sub
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    add_footer(s1, 1)

    # Common Card Grid Layout Parameters
    card_w = Inches(3.64)
    card_gap = Inches(0.4)
    grid_start_x = Inches(0.8)
    card_y = Inches(1.95)
    card_h = Inches(4.7)

    # =========================================================================
    # SLIDE 2: The Problem (Social Engineering Crisis)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "The Challenge & Threat Landscape", "Modern Scams Target Psychology, Not Software", "Legacy antivirus tools scan files for malware, but today's cybercriminals hack human trust.")

    cards_data_s2 = [
        (
            ACCENT_RED,
            "01",
            "$10B+ Annual Losses",
            "The Social Engineering Epidemic",
            [
                ("Psychological Exploitation", "Attackers leverage urgency, fear, authority, and financial panic rather than malware code."),
                ("Smishing & Impersonation", "Fraudulent bank alerts and fake parcel tracking SMS outpace traditional spam vectors."),
                ("Vulnerable Demographics", "Elderly and non-technical consumers bear the brunt of devastating financial fraud.")
            ]
        ),
        (
            ACCENT_AMBER,
            "02",
            "The Binary Alert Gap",
            "Legacy Antivirus is Opaque",
            [
                ("Unexplained Warnings", "Existing tools flag a message as 'suspicious' without explaining why it is fraudulent."),
                ("Confusion & Re-infection", "Without understanding the attack angle, victims often fall for the follow-up phone call or SMS."),
                ("Zero Actionable Guidance", "Users receive no prioritized recovery steps on how to freeze cards or revoke credentials.")
            ]
        ),
        (
            ACCENT_BLUE,
            "03",
            "Multi-Channel Infiltration",
            "Evolving Cross-Vector Scams",
            [
                ("Cross-Modal Camouflage", "Attacks start with an SMS, lead to a fake website, and finish on an encrypted messaging app."),
                ("Spoofed Visual Identities", "High-fidelity visual clones of Chase, PayPal, and USPS trick conventional email scanners."),
                ("Ephemeral Domains", "Malicious domains exist for only hours, slipping under static threat blacklist feeds.")
            ]
        )
    ]

    for i, (accent, num, stat, headline, bullets) in enumerate(cards_data_s2):
        cx = grid_start_x + i * (card_w + card_gap)
        c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)

        # Top Accent Line
        top_line = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = f"CRITICAL FACTOR #{num}"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent

        p_stat = tf.add_paragraph()
        p_stat.text = stat
        p_stat.font.name = "Segoe UI"
        p_stat.font.size = Pt(17)
        p_stat.font.bold = True
        p_stat.font.color.rgb = TEXT_WHITE

        p_h = tf.add_paragraph()
        p_h.text = headline
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(11.5)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_CYAN

        tf.add_paragraph().font.size = Pt(4)

        for b_title, b_desc in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b_title}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_WHITE

            run = pb.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    add_footer(s2, 2)

    # =========================================================================
    # SLIDE 3: The Solution (Explainable Threat Intelligence)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "The ScamGuard Paradigm", "Demystifying Cyber Threats in Plain English", "Replacing cryptic security errors with actionable, transparent forensic intelligence.")

    cards_data_s3 = [
        (
            ACCENT_BLUE,
            "PRECISION RISK SCORING",
            "Instant Threat Thermometer",
            "Calibrated 0-100% Risk Assessment",
            [
                ("Dynamic Scoring", "Calculated using composite weights from NLP intent, visual spoof indicators, and URL metadata."),
                ("Clear Risk Tiers", "Categorized into Critical (Red), Suspicious (Amber), or Safe (Green) for instant decision making."),
                ("Executive Verdict", "Plain-English diagnosis summarizing whether the message is phishing, extortion, or legitimate.")
            ]
        ),
        (
            ACCENT_CYAN,
            "ATTACK CHAIN DECONSTRUCTION",
            "Forensic Threat Mapping",
            "Visual Step-by-Step Methodology",
            [
                ("Infiltration Phase", "Highlights the hook: fake shipping updates, fraudulent account lockouts, or tax demands."),
                ("Exploitation Phase", "Exposes the trap: credential harvesting forms, session hijackers, or remote trojans."),
                ("Exfiltration Phase", "Details how the attacker profits: identity theft, unauthorized wire transfers, or SIM swapping.")
            ]
        ),
        (
            ACCENT_GREEN,
            "ACTIONABLE RECOVERY BLUEPRINT",
            "Personal Incident Response",
            "Prioritized Countermeasures",
            [
                ("Immediate Safeguards", "Clear instructions on what to do right now (e.g., 'Do NOT tap links', 'Freeze debit card')."),
                ("Account Hardening", "Guidance on enabling MFA, resetting primary passwords, and logging out active sessions."),
                ("Official Reporting", "Direct routing links to report abuse to carriers (7726), FTC, and brand fraud desks.")
            ]
        )
    ]

    for i, (accent, tag, title, sub, bullets) in enumerate(cards_data_s3):
        cx = grid_start_x + i * (card_w + card_gap)
        c_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)

        top_line = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = tag
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_CYAN

        tf.add_paragraph().font.size = Pt(4)

        for b_title, b_desc in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b_title}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_WHITE

            run = pb.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    add_footer(s3, 3)

    # =========================================================================
    # SLIDE 4: Multi-Modal Ingestion Engine
    # =========================================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "Multi-Modal Ingestion Architecture", "Three Seamless Modalities for Complete Defense", "Whether it's a suspicious text, an app screenshot, or a web link, ScamGuard analyzes it all.")

    modalities = [
        (
            ACCENT_BLUE,
            "MODALITY 01",
            "Message & Text Analysis",
            "Paste SMS, Emails, DMs, or Chats",
            [
                ("Linguistic Urgency Analysis", "Detects artificial panic triggers, coercive deadlines, and threatening extortion phrasing."),
                ("Impersonation Verification", "Cross-references claimed sender identity against authentic bank and carrier formats."),
                ("Unicode Sanitization", "Strips invisible zero-width characters, homoglyphs, and obfuscated text tricks.")
            ],
            "Fastest turnaround | Sub-second analysis"
        ),
        (
            ACCENT_CYAN,
            "MODALITY 02",
            "Screenshot & Image OCR",
            "Drop Photos of SMS, Screens & Mails",
            [
                ("High-Fidelity Vision OCR", "Extracts complex mobile conversation bubbles and embedded URL text via Gemini Vision."),
                ("Visual Deception Check", "Identifies counterfeit logos, spoofed interface layouts, and fake trust badges."),
                ("Layout Context Analysis", "Analyzes the overall visual composition to catch fake push notifications and login screens.")
            ],
            "Deep Vision AI | Supports PNG, JPG, HEIC"
        ),
        (
            ACCENT_GREEN,
            "MODALITY 03",
            "URL & Domain Deep-Scan",
            "Inspect Suspicious Links Safely",
            [
                ("Redirect Chain Tracking", "Safely traverses deceptive HTTP 301/302 redirects and shorteners without executing payload."),
                ("WHOIS & Domain Age", "Flags newly registered domains (< 30 days old) impersonating longstanding financial brands."),
                ("Threat Intelligence Query", "Cross-references global blacklist feeds and Safe Browsing intelligence databases.")
            ],
            "Zero-Click Inspection | Client-safe sandbox"
        )
    ]

    for i, (accent, tag, title, sub, bullets, footer_tag) in enumerate(modalities):
        cx = grid_start_x + i * (card_w + card_gap)
        c_shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)

        top_line = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = tag
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_CYAN

        tf.add_paragraph().font.size = Pt(4)

        for b_title, b_desc in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b_title}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_WHITE

            run = pb.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

        tf.add_paragraph().font.size = Pt(6)
        p_badge = tf.add_paragraph()
        p_badge.text = f"[ {footer_tag} ]"
        p_badge.font.name = "Segoe UI"
        p_badge.font.size = Pt(9)
        p_badge.font.bold = True
        p_badge.font.color.rgb = accent

    add_footer(s4, 4)

    # =========================================================================
    # SLIDE 5: Technical Architecture & Pipeline
    # =========================================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Full-Stack System Architecture", "High-Throughput Threat Intelligence Pipeline", "Built with modern decoupled architecture combining Next.js 16, FastAPI, and Google Gemini.")

    steps = [
        ("01", "Client Ingestion", "Next.js 16 / React 19", [
            "Tailwind CSS v4 design system",
            "Multi-modal input dispatcher",
            "Client-side payload sanitization",
            "Responsive dark/light UI"
        ], ACCENT_BLUE),
        ("02", "API Orchestrator", "FastAPI Asynchronous Backend", [
            "Pydantic strict schema validation",
            "Non-blocking async task dispatch",
            "Ephemeral sandbox coordinator",
            "CORS & rate-limiting guards"
        ], ACCENT_CYAN),
        ("03", "AI & Intel Core", "Gemini Engine + Threat Feeds", [
            "Multi-modal vision & text reasoning",
            "DNS resolver & redirect crawler",
            "WHOIS domain age inspection",
            "Known-bad scam template matching"
        ], ACCENT_AMBER),
        ("04", "Forensic Synthesis", "Structured Intelligence JSON", [
            "Deterministic JSON threat schema",
            "Weighted composite risk scoring",
            "Attack chain step generation",
            "Actionable user recommendations"
        ], ACCENT_GREEN)
    ]

    s5_card_w = Inches(2.72)
    s5_card_gap = Inches(0.28)
    s5_start_x = Inches(0.8)

    for i, (num, title, tech, points, accent) in enumerate(steps):
        cx = s5_start_x + i * (s5_card_w + s5_card_gap)
        c_shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, s5_card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)

        top_line = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, s5_card_w, Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = f"STAGE {num}"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.text = tech
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(10)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_CYAN

        tf.add_paragraph().font.size = Pt(6)

        for pt in points:
            pb = tf.add_paragraph()
            pb.text = f"✔ {pt}"
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = TEXT_MUTED

    add_footer(s5, 5)

    # =========================================================================
    # SLIDE 6: Product Experience (The Threat Report)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Product Interface & Output", "The Forensic Threat Report: Actionable, Not Overwhelming", "Transforming complex cyber intelligence into an intuitive, high-clarity dashboard.")

    features = [
        (
            ACCENT_RED,
            "PRIMARY VERDICT & SCORE",
            "Calibrated Risk Gauge",
            "Instant visual feedback for users",
            [
                ("High-Contrast Badge", "Prominent risk score (e.g., 92% CRITICAL) visible in milliseconds."),
                ("Executive Summary", "A clear 2-sentence summary of the fraudster's objective without technical jargon."),
                ("Category Classification", "Identifies specific threat type (Bank Impersonation, Credential Harvester, Extortion).")
            ]
        ),
        (
            ACCENT_CYAN,
            "FORENSIC DECONSTRUCTION",
            "Interactive Attack Chain",
            "Visualizing the cyber kill-chain",
            [
                ("Stage 1: The Lure", "Explains why the message was constructed to provoke urgency, panic, or greed."),
                ("Stage 2: The Mechanism", "Demonstrates how the link or attachment captures passwords or session cookies."),
                ("Stage 3: The Impact", "Illustrates the attacker's final monetization step to build lasting user vigilance.")
            ]
        ),
        (
            ACCENT_GREEN,
            "EMERGENCY PROTOCOLS",
            "Personal Action Checklist",
            "Step-by-step remediation guide",
            [
                ("Immediate Containment", "Tells the user exactly what NOT to touch and what connections to sever."),
                ("Account Recovery", "Targeted guidance on changing affected credentials and enabling hardware 2FA."),
                ("One-Click Reporting", "Built-in templates to submit complaints directly to carriers (7726) and federal agencies.")
            ]
        )
    ]

    for i, (accent, tag, title, sub, bullets) in enumerate(features):
        cx = grid_start_x + i * (card_w + card_gap)
        c_shape = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)

        top_line = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = tag
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_CYAN

        tf.add_paragraph().font.size = Pt(4)

        for b_title, b_desc in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b_title}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_WHITE

            run = pb.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    add_footer(s6, 6)

    # =========================================================================
    # SLIDE 7: Privacy & Security Architecture
    # =========================================================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Trust, Privacy & Governance", "Zero-Retention Architecture: Complete Privacy by Design", "Users submitting sensitive messages and screenshots must have absolute confidence in data privacy.")

    privacy_pillars = [
        (
            ACCENT_CYAN,
            "PILLAR 01",
            "Ephemeral Memory Execution",
            "Zero Data at Rest",
            [
                ("In-Memory Processing", "Incoming messages, images, and links are analyzed exclusively in transient RAM buffers."),
                ("Automatic Purging", "All processing memory is immediately released and destroyed once the JSON report is returned."),
                ("No Disk Caching", "No temporary files, image dumps, or log files are saved to server storage.")
            ]
        ),
        (
            ACCENT_GREEN,
            "PILLAR 02",
            "No Databases & No PII",
            "Zero User Profiling",
            [
                ("Zero Database Footprint", "There is no relational database, NoSQL store, or cloud bucket holding submitted scans."),
                ("No Account Required", "Users do not log in or provide email addresses, phone numbers, or identity tokens."),
                ("No Ad-Network Tracking", "Zero third-party marketing trackers, tracking pixels, or data-broker integrations.")
            ]
        ),
        (
            ACCENT_BLUE,
            "PILLAR 03",
            "End-to-End Encryption",
            "Hardened Transit Security",
            [
                ("TLS 1.3 Encryption", "All communications between client and server are secured using modern cryptographic suites."),
                ("Isolated Microservices", "Backend API services operate in sandboxed containers with strictly bounded network permissions."),
                ("Safe Link Crawler", "URL redirect resolution runs through isolated headless proxies with client sandboxing.")
            ]
        )
    ]

    for i, (accent, tag, title, sub, bullets) in enumerate(privacy_pillars):
        cx = grid_start_x + i * (card_w + card_gap)
        c_shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)

        top_line = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = accent
        top_line.line.fill.background()

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = tag
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_CYAN

        tf.add_paragraph().font.size = Pt(4)

        for b_title, b_desc in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b_title}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_WHITE

            run = pb.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    add_footer(s7, 7)

    # =========================================================================
    # SLIDE 8: Roadmap & Strategic Vision
    # =========================================================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "Future Roadmap & Strategic Vision", "The Future of Proactive Personal Scam Defense", "Scaling ScamGuard from on-demand web analysis to an omnipresent personal cybersecurity shield.")

    left_w = Inches(6.8)
    right_w = Inches(4.5)
    split_gap = Inches(0.433)

    # Left Card: Strategic Milestones
    left_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), card_y, left_w, card_h)
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = CARD_BG
    left_card.line.color.rgb = CARD_BORDER
    left_card.line.width = Pt(1)

    top_line_l = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), card_y, left_w, Inches(0.08))
    top_line_l.fill.solid()
    top_line_l.fill.fore_color.rgb = ACCENT_BLUE
    top_line_l.line.fill.background()

    tf_l = left_card.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = Inches(0.3)
    tf_l.margin_top = Inches(0.25)

    p0 = tf_l.paragraphs[0]
    p0.text = "PRODUCT EVOLUTION & MILESTONES"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE

    p_t = tf_l.add_paragraph()
    p_t.text = "Strategic Roadmap (2026 - 2027)"
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(17)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE

    tf_l.add_paragraph().font.size = Pt(4)

    milestones = [
        ("Phase 1: Browser Shield Extension", "Live on-page scanning of checkout forms, suspicious web redirects, and fake banking portals."),
        ("Phase 2: Mobile SMS Gateway", "Native iOS & Android integration for automatic, background threat screening of incoming SMS messages."),
        ("Phase 3: Crowdsourced Threat Radar", "Real-time anonymized telemetry to detect active localized scam campaigns before they spread."),
        ("Phase 4: Enterprise Fraud API", "Providing banking and fintech partners with automated fraud triage and incident response feeds.")
    ]

    for m_title, m_desc in milestones:
        pm = tf_l.add_paragraph()
        pm.text = f"▶ {m_title}\n"
        pm.font.name = "Segoe UI"
        pm.font.size = Pt(11)
        pm.font.bold = True
        pm.font.color.rgb = TEXT_CYAN

        run = pm.add_run()
        run.text = f"   {m_desc}"
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = TEXT_MUTED

    # Right Card: Mission & Impact Callout
    rx = Inches(0.8) + left_w + split_gap
    right_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, card_y, right_w, card_h)
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = RGBColor(14, 25, 48)
    right_card.line.color.rgb = ACCENT_BLUE
    right_card.line.width = Pt(1.5)

    tf_r = right_card.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    p_rtag = tf_r.paragraphs[0]
    p_rtag.text = "CORE VALUE PROPOSITION"
    p_rtag.font.name = "Segoe UI"
    p_rtag.font.size = Pt(9.5)
    p_rtag.font.bold = True
    p_rtag.font.color.rgb = ACCENT_GREEN

    p_rtitle = tf_r.add_paragraph()
    p_rtitle.text = "Empowering Everyone Against Cyber Crime"
    p_rtitle.font.name = "Segoe UI"
    p_rtitle.font.size = Pt(18)
    p_rtitle.font.bold = True
    p_rtitle.font.color.rgb = TEXT_WHITE

    tf_r.add_paragraph().font.size = Pt(8)

    p_quote = tf_r.add_paragraph()
    p_quote.text = "\"Cybersecurity is no longer just about firewalls and endpoint agents. In an era where attacks exploit human emotion, education and transparency are the ultimate armor.\""
    p_quote.font.name = "Segoe UI"
    p_quote.font.size = Pt(11.5)
    p_quote.font.italic = True
    p_quote.font.color.rgb = TEXT_CYAN

    tf_r.add_paragraph().font.size = Pt(8)

    impacts = [
        "100% Free & Accessible to Any Web User",
        "Sub-2-Second Multi-Modal AI Analysis",
        "Proven Defense Against Sophisticated Smishing",
        "Democratizing Enterprise-Grade Threat Intel"
    ]

    for imp in impacts:
        p_imp = tf_r.add_paragraph()
        p_imp.text = f"★  {imp}"
        p_imp.font.name = "Segoe UI"
        p_imp.font.size = Pt(10.5)
        p_imp.font.bold = True
        p_imp.font.color.rgb = TEXT_WHITE

    add_footer(s8, 8)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "ScamGuard_Presentation.pptx"
    build_presentation(out_file)
